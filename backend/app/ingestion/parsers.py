import io
import json
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from app.ingestion.detect import FileKind, guess_delimiter

NA_VALUES = ["", "NA", "N/A", "n/a", "null", "NULL", "None", "NaN", "nan"]


class ParseError(ValueError):
    pass


@dataclass
class ParsedTable:
    name_hint: str
    df: pd.DataFrame


@dataclass
class TextSegment:
    text: str
    page: int | None = None


@dataclass
class ParseResult:
    tables: list[ParsedTable] = field(default_factory=list)
    segments: list[TextSegment] = field(default_factory=list)


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _blank_to_none(df: pd.DataFrame) -> pd.DataFrame:
    return df.where(df != "", None)


def _parse_csv(path: Path, name: str) -> ParseResult:
    text = read_text(path)
    if not text.strip():
        raise ParseError("The file is empty.")
    try:
        df = pd.read_csv(
            io.StringIO(text),
            sep=guess_delimiter(text),
            dtype=str,
            keep_default_na=False,
            na_values=NA_VALUES,
            skipinitialspace=True,
        )
    except pd.errors.EmptyDataError as exc:
        raise ParseError("The file is empty.") from exc
    except pd.errors.ParserError as exc:
        raise ParseError(f"Could not parse the delimited file: {exc}") from exc
    return ParseResult(tables=[ParsedTable(Path(name).stem, df.dropna(how="all"))])


def _find_records(data):
    if isinstance(data, list) and data and sum(isinstance(r, dict) for r in data) / len(data) >= 0.8:
        return [r for r in data if isinstance(r, dict)]
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list) and v and all(isinstance(r, dict) for r in v)]
        if len(lists) == 1:
            return lists[0]
    return None


def _parse_json(path: Path, name: str) -> ParseResult:
    text = read_text(path).strip()
    if not text:
        raise ParseError("The file is empty.")
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        try:
            data = [json.loads(line) for line in text.splitlines() if line.strip()]
        except json.JSONDecodeError as exc:
            raise ParseError(f"Invalid JSON: {exc}") from exc
    records = _find_records(data)
    if records is None:
        return ParseResult(segments=[TextSegment(json.dumps(data, indent=2, ensure_ascii=False))])
    df = pd.json_normalize(records, sep="_")
    for col in df.columns:
        if df[col].map(lambda v: isinstance(v, (list, dict))).any():
            df[col] = df[col].map(lambda v: json.dumps(v) if isinstance(v, (list, dict)) else v)
    return ParseResult(tables=[ParsedTable(Path(name).stem, df)])


def _parse_excel(path: Path, name: str) -> ParseResult:
    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Could not read the Excel file: {exc}") from exc
    stem = Path(name).stem
    tables = []
    for sheet, df in sheets.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue
        tables.append(ParsedTable(stem if len(sheets) == 1 else f"{stem}_{sheet}", df))
    return ParseResult(tables=tables)


def _parse_pdf(path: Path, name: str) -> ParseResult:
    import pdfplumber

    stem = Path(name).stem
    result = ParseResult()
    try:
        with pdfplumber.open(path) as pdf:
            for pno, page in enumerate(pdf.pages, start=1):
                for k, raw in enumerate(page.extract_tables() or [], start=1):
                    rows = [[(c or "").strip() for c in r] for r in raw if r and any(c and c.strip() for c in r)]
                    if len(rows) < 2 or len(rows[0]) < 2:
                        continue
                    header = [h or f"col_{i + 1}" for i, h in enumerate(rows[0])]
                    body = [(r + [""] * len(header))[: len(header)] for r in rows[1:]]
                    df = _blank_to_none(pd.DataFrame(body, columns=header))
                    result.tables.append(ParsedTable(f"{stem}_table_p{pno}_{k}", df))
                text = page.extract_text() or ""
                if text.strip():
                    result.segments.append(TextSegment(text, page=pno))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Could not read the PDF: {exc}") from exc
    if not result.tables and not result.segments:
        raise ParseError("No extractable text found in the PDF (scanned/image-only PDFs are not supported).")
    return result


def _parse_docx(path: Path, name: str) -> ParseResult:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Could not read the Word document: {exc}") from exc
    stem = Path(name).stem
    result = ParseResult(segments=[TextSegment("\n".join(p.text for p in document.paragraphs))])
    for k, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) >= 2 and len(rows[0]) >= 2:
            result.tables.append(ParsedTable(f"{stem}_table_{k}", _blank_to_none(pd.DataFrame(rows[1:], columns=rows[0]))))
    return result


def _parse_text(path: Path, name: str) -> ParseResult:
    return ParseResult(segments=[TextSegment(read_text(path))])


PARSERS = {
    FileKind.CSV: _parse_csv,
    FileKind.JSON: _parse_json,
    FileKind.EXCEL: _parse_excel,
    FileKind.PDF: _parse_pdf,
    FileKind.DOCX: _parse_docx,
    FileKind.TEXT: _parse_text,
}


def parse_file(path: Path, kind: FileKind, original_name: str) -> ParseResult:
    result = PARSERS[kind](Path(path), original_name)
    result.tables = [t for t in result.tables if not t.df.empty and len(t.df.columns) > 0]
    result.segments = [s for s in result.segments if s.text.strip()]
    if not result.tables and not result.segments:
        raise ParseError("The file contains no readable data.")
    return result
